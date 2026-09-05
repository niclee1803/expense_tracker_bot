from collections import defaultdict
from io import BytesIO

from openpyxl import Workbook

from settings import TIMEZONE, category_name
from storage import get_all_expenses
from utils import now_local


def build_expenses_xlsx(user_id: int) -> BytesIO | None:
    expenses_by_month = defaultdict(list)

    for expense in get_all_expenses(user_id):
        target_month = expense.get("month")

        if target_month:
            expenses_by_month[target_month].append(expense)

    if not expenses_by_month:
        return None

    workbook = Workbook()
    workbook.remove(workbook.active)

    for target_month in sorted(expenses_by_month):
        _add_month_sheet(
            workbook,
            target_month,
            expenses_by_month[target_month],
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    output.name = (
        f"expense-data-{now_local().strftime('%Y-%m-%d')}.xlsx"
    )
    return output


def _add_month_sheet(
    workbook: Workbook,
    target_month: str,
    expenses: list[dict],
) -> None:
    worksheet = workbook.create_sheet(title=target_month)
    worksheet.append(["Date", "Category", "Amount", "Comment"])

    category_totals = defaultdict(int)

    for expense in expenses:
        created_at = expense.get("created_at")
        local_date = (
            created_at.astimezone(TIMEZONE).date()
            if created_at
            else None
        )

        category = expense.get("category", "others")
        amount_cents = int(expense["amount_cents"])
        comment = expense.get("comment", "")

        worksheet.append(
            [
                local_date,
                category_name(category),
                amount_cents / 100,
                comment,
            ]
        )

        category_totals[category] += amount_cents

    raw_data_end_row = worksheet.max_row

    # Leave columns E:F empty between the raw data and summary.
    worksheet["G1"] = "Category totals"
    worksheet["G2"] = "Category"
    worksheet["H2"] = "Amount"

    summary_row = 3

    for category, amount_cents in sorted(
        category_totals.items(),
        key=lambda item: item[1],
        reverse=True,
    ):
        worksheet.cell(
            row=summary_row,
            column=7,
            value=category_name(category),
        )
        worksheet.cell(
            row=summary_row,
            column=8,
            value=amount_cents / 100,
        )
        summary_row += 1

    worksheet.cell(
        row=summary_row,
        column=7,
        value="Total expenses for the month",
    )
    worksheet.cell(
        row=summary_row,
        column=8,
        value=sum(category_totals.values()) / 100,
    )

    for cell in worksheet["A"][1:]:
        cell.number_format = "yyyy-mm-dd"

    for cell in worksheet["C"][1:]:
        cell.number_format = "#,##0.00"

    for cell in worksheet["H"][1:summary_row]:
        cell.number_format = "#,##0.00"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:D{raw_data_end_row}"

    column_widths = {
        "A": 19,
        "B": 22,
        "C": 14,
        "D": 35,
        "E": 3,
        "F": 3,
        "G": 28,
        "H": 14,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width